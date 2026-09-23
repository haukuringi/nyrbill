"""
Breytir bilagogn.xlsx í JSON-skrár fyrir vefsíðuna.

Notkun:
    pip install openpyxl
    python xlsx_i_json.py bilagogn.xlsx

Býr til möppuna data/ með:
    data/bilar.json       – allar gerðir með útfærslum
    data/umbod.json       – umboðin
    data/stillingar.json  – upplýsingar um rafbílastyrk
"""

import json
import re
import sys
import unicodedata
from collections import OrderedDict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

# Dálkar sem verða að vera fylltir út
SKYLDA = ["id", "merki", "gerd", "utfaersla", "umbod", "flokkur",
          "arsgerd", "verd", "orkugjafi", "drif", "heimild", "uppfaert"]
TOLUR = ["arsgerd", "verd", "hestofl", "rafhlada_kwh", "draegni_km",
         "eydsla_l_100km", "saeti", "farangursrymi_l", "hrodun_0_100_s"]
HEILTOLUR = ["arsgerd", "verd", "hestofl", "draegni_km", "saeti", "farangursrymi_l"]
# Formúludálkar í töflureikninum – forritið reiknar þá sjálft í staðinn
HUNSA = ["styrkhaef", "verd_med_styrk"]
LOSUNARFRITT = ["rafmagn", "vetni"]

villur = []
vidvaranir = []


def slug(texti):
    """'Škoda Enyaq Þór' -> 'skoda-enyaq-thor'"""
    t = str(texti).lower()
    for a, b in {"þ": "th", "ð": "d", "æ": "ae", "ö": "o"}.items():
        t = t.replace(a, b)
    t = unicodedata.normalize("NFKD", t).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", t).strip("-")


def lesa_toflu(ws):
    """Les flipa þar sem fyrsta lína eru dálkaheiti. Skilar lista af dict."""
    raðir = list(ws.iter_rows(values_only=True))
    if not raðir:
        return []
    haus = [str(h).strip() if h is not None else None for h in raðir[0]]
    ut = []
    for nr, rod in enumerate(raðir[1:], start=2):
        if all(v is None or str(v).strip() == "" for v in rod):
            continue
        d = {}
        for h, v in zip(haus, rod):
            if h is None:
                continue
            if isinstance(v, str):
                v = v.strip() or None
            d[h] = v
        d["_lina"] = nr
        ut.append(d)
    return ut


def hreinsa_tolu(gildi, dalkur, lina):
    if gildi is None:
        return None
    if isinstance(gildi, (int, float)):
        tala = gildi
    else:
        # Leyfir t.d. "5.990.000 kr." eða "7,9" ef einhver gleymdi sér
        s = str(gildi).replace("kr.", "").replace(" ", "")
        if dalkur == "verd":
            s = s.replace(".", "")
        s = s.replace(",", ".")
        try:
            tala = float(s)
        except ValueError:
            villur.append(f"Lína {lina}: '{dalkur}' á að vera tala, fékk '{gildi}'")
            return None
        vidvaranir.append(f"Lína {lina}: '{dalkur}' var texti ('{gildi}'), túlkað sem {tala}")
    return int(round(tala)) if dalkur in HEILTOLUR else float(tala)


def hreinsa_dags(gildi, lina):
    if gildi is None:
        return None
    if isinstance(gildi, (datetime, date)):
        return gildi.strftime("%Y-%m-%d")
    s = str(gildi).strip()
    for snid in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, snid).strftime("%Y-%m-%d")
        except ValueError:
            pass
    villur.append(f"Lína {lina}: 'uppfaert' er ekki gild dagsetning ('{gildi}')")
    return None


def lesa_stillingar(wb):
    if "Stillingar" not in wb.sheetnames:
        vidvaranir.append("Enginn Stillingar-flipi – nota sjálfgefin gildi fyrir styrk")
        return 10_000_000, 500_000
    ws = wb["Stillingar"]
    return int(ws["B2"].value), int(ws["B3"].value)


def main():
    if len(sys.argv) < 2:
        print("Notkun: python xlsx_i_json.py bilagogn.xlsx [utmappa]")
        sys.exit(1)
    skra = Path(sys.argv[1])
    utmappa = Path(sys.argv[2]) if len(sys.argv) > 2 else Path("data")

    wb = load_workbook(skra, data_only=True)
    verdthak, styrkur = lesa_stillingar(wb)

    # --- Leyfð gildi
    leyfd = {}
    if "Listar" in wb.sheetnames:
        ws = wb["Listar"]
        for col in ws.iter_cols(values_only=True):
            if col[0]:
                leyfd[col[0]] = {v for v in col[1:] if v}

    # --- Umboð
    umbod = []
    for r in lesa_toflu(wb["Umboð"]):
        if not r.get("id"):
            villur.append(f"Umboð lína {r['_lina']}: vantar id")
            continue
        umbod.append({"id": r["id"], "nafn": r.get("nafn"), "vefsida": r.get("vefsida")})
    umbod_ids = {u["id"] for u in umbod}

    # --- Útfærslur
    gerdir = OrderedDict()
    sed_ids = set()
    for r in lesa_toflu(wb["Útfærslur"]):
        lina = r["_lina"]

        if r.get("athugasemd") and "DÆMI" in str(r["athugasemd"]).upper():
            vidvaranir.append(f"Lína {lina}: sleppt, merkt sem DÆMI")
            continue

        vantar = [d for d in SKYLDA if r.get(d) is None]
        if vantar:
            villur.append(f"Lína {lina}: vantar {', '.join(vantar)}")
            continue

        if r["id"] in sed_ids:
            villur.append(f"Lína {lina}: id '{r['id']}' er notað oftar en einu sinni")
            continue
        sed_ids.add(r["id"])

        for d in ("flokkur", "orkugjafi", "drif"):
            if d in leyfd and r[d] not in leyfd[d]:
                villur.append(f"Lína {lina}: '{r[d]}' er ekki leyfilegt gildi fyrir {d}")
        if r["umbod"] not in umbod_ids:
            villur.append(f"Lína {lina}: umboðið '{r['umbod']}' er ekki í Umboð-flipanum")

        for d in TOLUR:
            r[d] = hreinsa_tolu(r.get(d), d, lina)
        r["uppfaert"] = hreinsa_dags(r["uppfaert"], lina)

        if r["verd"] is not None and not (500_000 <= r["verd"] <= 100_000_000):
            vidvaranir.append(f"Lína {lina}: verðið {r['verd']:,} kr. virðist óeðlilegt")

        styrkhaef = r["orkugjafi"] in LOSUNARFRITT and (r["verd"] or 0) < verdthak

        utfaersla = {
            "id": r["id"],
            "heiti": r["utfaersla"],
            "verd": r["verd"],
            "styrkhaef": styrkhaef,
            "verd_med_styrk": r["verd"] - styrkur if styrkhaef else r["verd"],
            "orkugjafi": r["orkugjafi"],
            "drif": r["drif"],
            "hestofl": r.get("hestofl"),
            "rafhlada_kwh": r.get("rafhlada_kwh"),
            "draegni_km": r.get("draegni_km"),
            "eydsla_l_100km": r.get("eydsla_l_100km"),
            "saeti": r.get("saeti"),
            "farangursrymi_l": r.get("farangursrymi_l"),
            "hrodun_0_100_s": r.get("hrodun_0_100_s"),
            "heimild": r["heimild"],
            "uppfaert": r["uppfaert"],
        }

        lykill = (r["merki"], r["gerd"], r["arsgerd"])
        if lykill not in gerdir:
            gerdir[lykill] = {
                "id": slug(f"{r['merki']}-{r['gerd']}-{r['arsgerd']}"),
                "merki": r["merki"],
                "gerd": r["gerd"],
                "umbod": r["umbod"],
                "flokkur": r["flokkur"],
                "arsgerd": r["arsgerd"],
                "utfaerslur": [],
            }
        gerdir[lykill]["utfaerslur"].append(utfaersla)

    # Lægsta verð og nýjasta uppfærsla á hverja gerð – hentugt fyrir listasíðuna
    bilar = []
    for g in gerdir.values():
        g["utfaerslur"].sort(key=lambda u: u["verd"])
        g["verd_fra"] = g["utfaerslur"][0]["verd"]
        g["verd_med_styrk_fra"] = min(u["verd_med_styrk"] for u in g["utfaerslur"])
        g["uppfaert"] = max(u["uppfaert"] for u in g["utfaerslur"] if u["uppfaert"])
        bilar.append(g)
    bilar.sort(key=lambda g: (g["merki"], g["gerd"]))

    # --- Niðurstöður
    for v in vidvaranir:
        print(f"  Aðvörun: {v}")
    if villur:
        print(f"\n{len(villur)} villur fundust – engar skrár skrifaðar:")
        for v in villur:
            print(f"  ✗ {v}")
        sys.exit(1)

    utmappa.mkdir(parents=True, exist_ok=True)

    def skrifa(nafn, gogn):
        with open(utmappa / nafn, "w", encoding="utf-8") as f:
            json.dump(gogn, f, ensure_ascii=False, indent=2)

    skrifa("bilar.json", bilar)
    skrifa("umbod.json", umbod)
    skrifa("stillingar.json", {
        "rafbilastyrkur": {"upphaed": styrkur, "verdthak": verdthak,
                           "heimild": "https://island.is/rafbilastyrkir"},
        "buid_til": datetime.now().strftime("%Y-%m-%d %H:%M"),
    })

    fj_utf = sum(len(g["utfaerslur"]) for g in bilar)
    print(f"\n✓ {len(bilar)} gerðir, {fj_utf} útfærslur, {len(umbod)} umboð → {utmappa}/")


if __name__ == "__main__":
    main()
